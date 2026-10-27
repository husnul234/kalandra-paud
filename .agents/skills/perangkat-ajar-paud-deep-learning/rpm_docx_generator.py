#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generator Dokumen DOCX Resmi: Perencanaan Pembelajaran Mendalam (RPM) PAUD Fase Fondasi
Standar Resmi Lulus Verifikasi Pengawas Kemendikdasmen RI
"""

import os
import sys
import json
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Palette Warna Resmi (Navy Deep Ocean)
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x8A)     # #1E3A8A Primary Navy
COLOR_HEADER_BLUE = RGBColor(0x1E, 0x40, 0xAF) # #1E40AF
COLOR_DARK = RGBColor(0x0F, 0x17, 0x2A)    # #0F172A Slate 900
COLOR_BODY = RGBColor(0x33, 0x41, 0x55)    # #334155 Slate 700
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)   # #64748B Slate 500

HEX_NAVY_BG = "1E3A8A"
HEX_SUBHEADER_BG = "F1F5F9"
HEX_ZEBRA_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_BORDER_DARK = "94A3B8"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color=HEX_BORDER, sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_cell_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def create_rpm_docx(data: dict, output_filepath: str):
    doc = docx.Document()

    # 1. Page Setup (Standar Administrasi Resmi RI: Kiri 3cm, Atas 2.5cm, Kanan 2cm, Bawah 2.5cm)
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.page_width = Cm(21.0)  # A4 Width
    section.page_height = Cm(29.7) # A4 Height

    info = data.get("informasi_umum", {})
    ident = data.get("identifikasi", {})
    desain = data.get("desain_pembelajaran", {})
    pengalaman = data.get("pengalaman_belajar", {})
    asesmen = data.get("asesmen_pembelajaran", {})
    iktp = data.get("iktp_bergradasi", [])
    rubrik = data.get("rubrik_asesmen", [])
    keselamatan = data.get("catatan_keselamatan", [])
    lampiran = data.get("lampiran_instrumen", {})

    school = info.get("asal_sekolah", "Satuan PAUD").upper()
    topic = info.get("topik_subtopik", "Pembelajaran PAUD")
    guru = info.get("guru", "Guru PAUD")
    kelas = info.get("jenjang_kelas", "Kelompok A (4-5 Tahun)")

    # Header Dokumen Resmi
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_h1 = p_header.add_run("PERENCANAAN PEMBELAJARAN MENDALAM (RPM)")
    run_h1.bold = True
    run_h1.font.name = "Calibri"
    run_h1.font.size = Pt(14)
    run_h1.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_h2 = p_sub.add_run(f"{school} — FASE FONDASI")
    run_h2.bold = True
    run_h2.font.name = "Calibri"
    run_h2.font.size = Pt(12)
    run_h2.font.color.rgb = COLOR_HEADER_BLUE

    p_year = doc.add_paragraph()
    p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_y = p_year.add_run("TAHUN AJARAN 2026/2027")
    run_y.font.name = "Calibri"
    run_y.font.size = Pt(10)
    run_y.font.color.rgb = COLOR_MUTED

    # Metadata Grid Table
    t_info = doc.add_table(rows=1, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_info, color=HEX_BORDER_DARK, sz="6")
    make_cell_cant_split(t_info.rows[0])

    c0 = t_info.cell(0, 0)
    set_cell_background(c0, HEX_SUBHEADER_BG)
    set_cell_margins(c0, 100, 100, 140, 140)
    c0.paragraphs[0].add_run(f"Nama Guru : {guru}\nFase / Kelas : Fondasi / {kelas}\nAlokasi Waktu : {info.get('alokasi_waktu', '6 x 3 JP')}").font.size = Pt(9.5)

    c1 = t_info.cell(0, 1)
    set_cell_background(c1, HEX_SUBHEADER_BG)
    set_cell_margins(c1, 100, 100, 140, 140)
    c1.paragraphs[0].add_run(f"Semester / Bln : {info.get('semester', '1 (Gasal)')} / {info.get('bulan', 'Juli 2026')}\nMinggu Ke- : Minggu ke-{info.get('minggu_ke', '1')}\nModel / Rombel : {info.get('model_pembelajaran', 'Sentra')} ({info.get('jumlah_anak', '15')} Anak)").font.size = Pt(9.5)

    doc.add_paragraph() # Spacer

    # Helper Section Title
    def add_section_header(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_NAVY

    # TABEL 1: IDENTIFIKASI
    add_section_header("I. IDENTIFIKASI PEMBELAJARAN")
    t1 = doc.add_table(rows=3, cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)
    
    t1.cell(0, 0).paragraphs[0].add_run("Karakteristik Peserta Didik").bold = True
    t1.cell(0, 1).paragraphs[0].add_run(ident.get("peserta_didik", "-"))
    
    t1.cell(1, 0).paragraphs[0].add_run("Materi Pelajaran").bold = True
    t1.cell(1, 1).paragraphs[0].add_run(ident.get("materi_pelajaran", "-"))
    
    t1.cell(2, 0).paragraphs[0].add_run("Dimensi Profil Lulusan (DPL)").bold = True
    dpl_text = ", ".join(ident.get("dimensi_profil_lulusan_terpilih", []))
    t1.cell(2, 1).paragraphs[0].add_run(f"Target DPL: {dpl_text}")

    for row in t1.rows:
        make_cell_cant_split(row)
        set_cell_background(row.cells[0], HEX_SUBHEADER_BG)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        row.cells[0].width = Cm(5.0)
        row.cells[1].width = Cm(11.0)

    doc.add_paragraph()

    # TABEL 2: DESAIN PEMBELAJARAN
    add_section_header("II. DESAIN PEMBELAJARAN")
    t2 = doc.add_table(rows=6, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    fields_t2 = [
        ("Capaian Pembelajaran (CP Verbatim)", "\n\n".join([f"• {c}" for c in desain.get("capaian_pembelajaran", [])])),
        ("Lintas Disiplin Ilmu", desain.get("lintas_disiplin_ilmu", "-")),
        ("Tujuan Pembelajaran", "\n".join([f"{idx+1}. {tp}" for idx, tp in enumerate(desain.get("tujuan_pembelajaran", []))])),
        ("Topik / Subtopik", desain.get("topik_pembelajaran", topic)),
        ("Praktik Pedagogis (Deep Learning)", desain.get("praktik_pedagogis", "-")),
        ("Pemanfaatan Digital", f"Perencanaan: {desain.get('pemanfaatan_digital', {}).get('perencanaan', '-')}\nPelaksanaan: {desain.get('pemanfaatan_digital', {}).get('pelaksanaan', '-')}\nAsesmen: {desain.get('pemanfaatan_digital', {}).get('asesmen', '-')}")
    ]

    for idx, (label, val) in enumerate(fields_t2):
        row = t2.rows[idx]
        make_cell_cant_split(row)
        set_cell_background(row.cells[0], HEX_SUBHEADER_BG)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(val)
        row.cells[0].width = Cm(5.0)
        row.cells[1].width = Cm(11.0)

    doc.add_paragraph()

    # TABEL 3: PENGALAMAN BELAJAR 6 HARI (18 KEGIATAN MAIN)
    add_section_header("III. PENGALAMAN BELAJAR MENDALAM (SIKLUS 6 HARI)")
    
    # Tahap Awal
    p_awal = doc.add_paragraph()
    p_awal.add_run("A. Tahap Awal (Berkesadaran, Bermakna, Menggembirakan):").bold = True
    for item in pengalaman.get("awal", {}).get("poin_kegiatan", []):
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(item)

    # Inti 6 Hari Table
    doc.add_paragraph().add_run("B. Tahap Inti (Memahami, Mengaplikasi, Merefleksi — 18 Kegiatan Main Loose Parts):").bold = True
    
    t3 = doc.add_table(rows=1, cols=4)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)
    
    # Table Header
    hdr = t3.rows[0]
    set_cell_background(hdr.cells[0], HEX_NAVY_BG)
    set_cell_background(hdr.cells[1], HEX_NAVY_BG)
    set_cell_background(hdr.cells[2], HEX_NAVY_BG)
    set_cell_background(hdr.cells[3], HEX_NAVY_BG)
    
    hdr_titles = ["Hari / Tahap", "No & Nama Kegiatan Main", "Alat dan Bahan (Loose Parts)", "Langkah Pelaksanaan & Mitigasi"]
    for i, title in enumerate(hdr_titles):
        set_cell_margins(hdr.cells[i], 120, 120, 100, 100)
        r = hdr.cells[i].paragraphs[0].add_run(title)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    blok_list = pengalaman.get("inti", {}).get("blok_tahap", [])
    for blok in blok_list:
        header_tahap = blok.get("header_tahap", "")
        for hari in blok.get("hari", []):
            hari_ke = hari.get("hari_ke", 1)
            kegiatan_list = hari.get("kegiatan_list", [])
            for k in kegiatan_list:
                row = t3.add_row()
                make_cell_cant_split(row)
                set_cell_margins(row.cells[0])
                set_cell_margins(row.cells[1])
                set_cell_margins(row.cells[2])
                set_cell_margins(row.cells[3])

                row.cells[0].paragraphs[0].add_run(f"Hari ke-{hari_ke}\n({header_tahap.split('—')[0].strip()})").font.size = Pt(9)
                row.cells[1].paragraphs[0].add_run(f"Kegiatan {k.get('nomor', 1)}:\n{k.get('nama_kegiatan', '-')}\n[{k.get('label', '-')}]").font.size = Pt(9)
                row.cells[2].paragraphs[0].add_run(k.get("alat_dan_bahan", "-")).font.size = Pt(9)
                
                langkah_str = "\n".join([f"- {l}" for l in k.get("langkah", [])])
                safety_str = f"\n⚠️ Catatan: {', '.join(k.get('catatan_keselamatan', []))}" if k.get("catatan_keselamatan") else ""
                row.cells[3].paragraphs[0].add_run(f"{k.get('judul_langkah', 'Cara Bermain')}:\n{langkah_str}{safety_str}").font.size = Pt(9)

    doc.add_paragraph()

    # Tahap Penutup
    p_penutup = doc.add_paragraph()
    p_penutup.add_run("C. Tahap Penutup (Berkesadaran, Menggembirakan):").bold = True
    for item in pengalaman.get("penutup", {}).get("poin_kegiatan", []):
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(item)

    doc.add_paragraph()

    # TABEL 4: ASESMEN PEMBELAJARAN
    add_section_header("IV. ASESMEN PEMBELAJARAN")
    t4 = doc.add_table(rows=3, cols=2)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t4)
    
    asesmen_rows = [
        ("Asesmen Awal (Diagnostik)", "\n".join([f"• {a}" for a in asesmen.get("awal", [])])),
        ("Asesmen Proses (Formatif)", "\n".join([f"• {a}" for a in asesmen.get("proses", [])])),
        ("Asesmen Akhir (Sumatif)", "\n".join([f"• {a}" for a in asesmen.get("akhir", [])]))
    ]
    for idx, (lbl, val) in enumerate(asesmen_rows):
        row = t4.rows[idx]
        make_cell_cant_split(row)
        set_cell_background(row.cells[0], HEX_SUBHEADER_BG)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        row.cells[0].paragraphs[0].add_run(lbl).bold = True
        row.cells[1].paragraphs[0].add_run(val)
        row.cells[0].width = Cm(5.0)
        row.cells[1].width = Cm(11.0)

    doc.add_paragraph()

    # TABEL 5: IKTP BERGRADASI SOLO
    add_section_header("V. INDIKATOR KETERCAPAIAN TUJUAN PEMBELAJARAN (IKTP) BERGRADASI TAKSONOMI SOLO")
    t5 = doc.add_table(rows=1, cols=5)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t5)
    
    hdr5 = t5.rows[0]
    titles5 = ["Tujuan Pembelajaran", "BB (Unistruktural)", "MB (Multistruktural)", "BSH (Relasional)", "BSB (Extended Abstract)"]
    for i, t in enumerate(titles5):
        set_cell_background(hdr5.cells[i], HEX_NAVY_BG)
        set_cell_margins(hdr5.cells[i], 120, 120, 80, 80)
        r = hdr5.cells[i].paragraphs[0].add_run(t)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(8.5)

    for item in iktp:
        row = t5.add_row()
        make_cell_cant_split(row)
        for i in range(5):
            set_cell_margins(row.cells[i], 100, 100, 80, 80)
        row.cells[0].paragraphs[0].add_run(f"TP {item.get('tujuan_ke', 1)}:\n{item.get('tujuan_ringkas', '-')}")
        row.cells[1].paragraphs[0].add_run(item.get("solo_bb", "-"))
        row.cells[2].paragraphs[0].add_run(item.get("solo_mb", "-"))
        row.cells[3].paragraphs[0].add_run(item.get("solo_bsh", "-"))
        row.cells[4].paragraphs[0].add_run(item.get("solo_bsb", "-"))

    doc.add_paragraph()

    # TABEL 6: RUBRIK ASESMEN 4 SKALA
    add_section_header("VI. RUBRIK ASESMEN 4 SKALA")
    t6 = doc.add_table(rows=1, cols=5)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t6)
    
    hdr6 = t6.rows[0]
    titles6 = ["Kriteria Asesmen", "Belum Berkembang (BB)", "Mulai Berkembang (MB)", "Berkembang Sesuai Harapan (BSH)", "Berkembang Sangat Baik (BSB)"]
    for i, t in enumerate(titles6):
        set_cell_background(hdr6.cells[i], HEX_NAVY_BG)
        set_cell_margins(hdr6.cells[i], 120, 120, 80, 80)
        r = hdr6.cells[i].paragraphs[0].add_run(t)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(8.5)

    for item in rubrik:
        row = t6.add_row()
        make_cell_cant_split(row)
        for i in range(5):
            set_cell_margins(row.cells[i], 100, 100, 80, 80)
        row.cells[0].paragraphs[0].add_run(item.get("kriteria", "-")).bold = True
        row.cells[1].paragraphs[0].add_run(item.get("bb", "-"))
        row.cells[2].paragraphs[0].add_run(item.get("mb", "-"))
        row.cells[3].paragraphs[0].add_run(item.get("bsh", "-"))
        row.cells[4].paragraphs[0].add_run(item.get("bsb", "-"))

    doc.add_paragraph()

    # TABEL 7: CATATAN KESELAMATAN
    if keselamatan:
        add_section_header("VII. CATATAN KESELAMATAN ANAK (RISK MITIGATION)")
        t7 = doc.add_table(rows=len(keselamatan), cols=1)
        t7.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t7, color=HEX_BORDER)
        for i, note in enumerate(keselamatan):
            row = t7.rows[i]
            make_cell_cant_split(row)
            set_cell_margins(row.cells[0])
            set_cell_background(row.cells[0], HEX_ZEBRA_BG)
            row.cells[0].paragraphs[0].add_run(f"⚠️ Prosedur {i+1}: {note}").font.size = Pt(9.5)

    doc.add_paragraph()

    # LEMBAR TANDA TANGAN
    p_ttd_loc = doc.add_paragraph()
    p_ttd_loc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd_loc.add_run(f"{info.get('asal_sekolah', 'Satuan PAUD')}, 20 Juli 2026\n\n")

    t_ttd = doc.add_table(rows=1, cols=2)
    t_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    # no borders
    for row in t_ttd.rows:
        make_cell_cant_split(row)
        set_cell_margins(row.cells[0], 0, 0, 0, 0)
        set_cell_margins(row.cells[1], 0, 0, 0, 0)
        row.cells[0].paragraphs[0].add_run("Mengetahui,\nKepala Satuan PAUD\n\n\n\n( .................................................... )\nNIP. .................................................").font.size = Pt(10)
        row.cells[1].paragraphs[0].add_run(f"Guru Kelas,\n{guru}\n\n\n\n( {guru} )\nNIP. .................................................").font.size = Pt(10)

    # 4 LAMPIRAN (Page Break)
    doc.add_page_break()

    p_lamp_hdr = doc.add_paragraph()
    p_lamp_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l1 = p_lamp_hdr.add_run("LAMPIRAN INSTRUMEN OBSERVASI AUTENTIK KELAS")
    r_l1.bold = True
    r_l1.font.size = Pt(13)
    r_l1.font.color.rgb = COLOR_NAVY

    doc.add_paragraph().add_run(f"Topik: {topic} | Fase Fondasi ({kelas})").bold = True

    lamp_items = [
        ("Lampiran 1: Format & Contoh Catatan Anekdot", lampiran.get("catatan_anekdot", "-")),
        ("Lampiran 2: Format Matriks Ceklis IKTP Per Anak", lampiran.get("ceklis_iktp", "-")),
        ("Lampiran 3: Format Penilaian Dokumentasi Hasil Karya", lampiran.get("hasil_karya", "-")),
        ("Lampiran 4: Format Portofolio Visual Foto Berseri", lampiran.get("foto_berseri", "-")),
    ]

    for title_l, content_l in lamp_items:
        add_section_header(title_l)
        p_l = doc.add_paragraph()
        p_l.paragraph_format.space_after = Pt(8)
        p_l.add_run(content_l)

    doc.save(output_filepath)
    print(f"File DOCX berhasil dibuat: {output_filepath}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        json_file = sys.argv[1]
        out_file = sys.argv[2] if len(sys.argv) > 2 else "RPM_PAUD_Fase_Fondasi.docx"
        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        create_rpm_docx(data, out_file)
    else:
        print("Penggunaan: python python_docx_engine.py data.json output.docx")
